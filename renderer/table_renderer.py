"""
Shared Table Renderer
Renders native editable PowerPoint tables styled with header fills, alternating row colors,
precise column widths, and row heights. Shared by 04, 06, and 08.
"""

from pptx.slide import Slide as PptxSlide
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from layout.coordinate_system import bbox_to_pptx_emu, px_to_emu
from layout.table_layout import TableLayoutResult
from config.theme import Theme, hex_to_rgb_color


class TableRenderer:
    @staticmethod
    def render_table(
        slide: PptxSlide,
        layout_result: TableLayoutResult,
        x: float,
        y: float,
        width: float,
        height: float
    ):
        """
        Renders a native PowerPoint table shape from TableLayoutResult.
        """
        num_rows = len(layout_result.rows)
        num_cols = len(layout_result.column_widths)
        if num_rows == 0 or num_cols == 0:
            return None

        left, top, w, h = bbox_to_pptx_emu(x, y, width, height)
        table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, w, h)
        table = table_shape.table

        # 1. Apply column widths
        for col_idx, col_w in enumerate(layout_result.column_widths):
            table.columns[col_idx].width = px_to_emu(col_w)

        # 2. Apply rows and cell contents
        for row_idx, row_layout in enumerate(layout_result.rows):
            # Row height
            table.rows[row_idx].height = px_to_emu(row_layout.allocated_height)
            is_header = row_layout.is_header

            for col_idx, cell_layout in enumerate(row_layout.cells):
                cell = table.cell(row_idx, col_idx)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE

                # Cell padding
                if num_rows >= 8:
                    cell.margin_left = Pt(3)
                    cell.margin_right = Pt(3)
                    cell.margin_top = Pt(1.5)
                    cell.margin_bottom = Pt(1.5)
                else:
                    cell.margin_left = Pt(4)
                    cell.margin_right = Pt(4)
                    cell.margin_top = Pt(2)
                    cell.margin_bottom = Pt(2)

                # Fill color
                cell.fill.solid()
                if is_header:
                    cell.fill.fore_color.rgb = hex_to_rgb_color(Theme.PRIMARY_NAVY)
                else:
                    row_bg = Theme.ROW_ALT_FILL if (row_idx % 2 == 1) else Theme.BG_WHITE
                    cell.fill.fore_color.rgb = hex_to_rgb_color(row_bg)

                # Cell text
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = cell_layout.text
                p.font.name = Theme.FONT_PRIMARY
                p.font.size = Pt(cell_layout.font_size)
                p.font.bold = is_header
                p.font.color.rgb = hex_to_rgb_color(Theme.TEXT_WHITE if is_header else Theme.TEXT_DARK)

                if cell_layout.align == "center":
                    p.alignment = PP_ALIGN.CENTER
                elif cell_layout.align == "right":
                    p.alignment = PP_ALIGN.RIGHT
                else:
                    p.alignment = PP_ALIGN.LEFT

        return table_shape
