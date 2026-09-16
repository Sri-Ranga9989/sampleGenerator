"""
Shared List Renderer
Renders diamond-badge TOC items (02_toc_image), section index items (03_section_opener),
and hierarchical bullet lists.
"""

from pptx.slide import Slide as PptxSlide
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
from layout.coordinate_system import bbox_to_pptx_emu, px_to_pt
from config.theme import Theme, hex_to_rgb_color
from models.report_model import TOCItem, IndexItem, BulletListBlock, BulletItem


class ListRenderer:
    @staticmethod
    def render_toc_item(
        slide: PptxSlide,
        item: TOCItem,
        x: float,
        y: float,
        width: float,
        height: float,
        font_size: float = 14.0,
        badge_size: float = 52.0
    ):
        """
        Renders a canonical toc_item:
        - Diamond shape badge with centered white number
        - Wrapping title text to the right of the badge
        """
        badge_left, badge_top, b_w, b_h = bbox_to_pptx_emu(x, y, badge_size, badge_size)
        diamond = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, badge_left, badge_top, b_w, b_h)
        diamond.fill.solid()
        diamond.fill.fore_color.rgb = hex_to_rgb_color(Theme.PRIMARY_NAVY)
        diamond.line.fill.background()

        # Text inside diamond
        tf_d = diamond.text_frame
        tf_d.word_wrap = False
        tf_d.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_d = tf_d.paragraphs[0]
        p_d.text = item.number
        p_d.font.name = Theme.FONT_PRIMARY
        p_d.font.size = Pt(font_size)
        p_d.font.bold = True
        p_d.font.color.rgb = hex_to_rgb_color(Theme.TEXT_WHITE)
        p_d.alignment = PP_ALIGN.CENTER

        # Label text box
        label_x = x + badge_size + 18.0
        label_w = max(20.0, width - badge_size - 22.0)
        lbl_left, lbl_top, l_w, l_h = bbox_to_pptx_emu(label_x, y + 6.0, label_w, height)
        tx_box = slide.shapes.add_textbox(lbl_left, lbl_top, l_w, l_h)
        tf_l = tx_box.text_frame
        tf_l.word_wrap = True
        tf_l.margin_left = 0
        tf_l.margin_right = 0
        tf_l.margin_top = 0
        tf_l.margin_bottom = 0

        p_l = tf_l.paragraphs[0]
        p_l.text = item.title
        p_l.font.name = Theme.FONT_PRIMARY
        p_l.font.size = Pt(font_size)
        p_l.font.bold = True
        p_l.font.color.rgb = hex_to_rgb_color(Theme.TEXT_DARK)
        p_l.alignment = PP_ALIGN.LEFT

        return diamond, tx_box

    @staticmethod
    def render_index_item(
        slide: PptxSlide,
        item: IndexItem,
        x: float,
        y: float,
        width: float,
        height: float,
        font_size: float = 13.0
    ):
        """
        Renders a canonical index_item (subsection index):
        - Bold subsection number (e.g. '2.1')
        - Subsection title text
        """
        left, top, w, h = bbox_to_pptx_emu(x, y, width, height)
        tx_box = slide.shapes.add_textbox(left, top, w, h)
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_top = 0

        p = tf.paragraphs[0]
        # Number run
        r_num = p.add_run()
        r_num.text = f"{item.number}  "
        r_num.font.name = Theme.FONT_PRIMARY
        r_num.font.size = Pt(font_size)
        r_num.font.bold = True
        r_num.font.color.rgb = hex_to_rgb_color(Theme.PRIMARY_NAVY)

        # Title run
        r_title = p.add_run()
        r_title.text = item.title
        r_title.font.name = Theme.FONT_PRIMARY
        r_title.font.size = Pt(font_size)
        r_title.font.bold = False
        r_title.font.color.rgb = hex_to_rgb_color(Theme.TEXT_DARK)

        return tx_box

    @staticmethod
    def render_bullet_list(
        slide: PptxSlide,
        bullet_list: BulletListBlock,
        x: float,
        y: float,
        width: float,
        height: float,
        font_size: float = 11.0
    ):
        """Renders hierarchical bullet list with custom indentation."""
        left, top, w, h = bbox_to_pptx_emu(x, y, width, height)
        tx_box = slide.shapes.add_textbox(left, top, w, h)
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(2)
        tf.margin_top = Pt(2)

        for idx, item in enumerate(bullet_list.items):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"•  {item.text}"
            p.level = item.level
            p.font.name = Theme.FONT_PRIMARY
            p.font.size = Pt(font_size)
            p.font.color.rgb = hex_to_rgb_color(Theme.TEXT_DARK)
            p.space_after = Pt(4)

        return tx_box
