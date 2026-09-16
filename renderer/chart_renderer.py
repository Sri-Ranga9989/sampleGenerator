"""
Shared Chart Renderer
Renders native editable PowerPoint charts (Bar, Line, Pie) with unified theme styling.
Shared across 04_table_chart and 07_large_chart.
"""

from typing import List, Optional
from pptx.slide import Slide as PptxSlide
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Pt
from layout.coordinate_system import bbox_to_pptx_emu
from layout.chart_layout import ChartLayoutResult
from config.theme import Theme, hex_to_rgb_color
from models.report_model import ChartBlock


class ChartRenderer:
    @staticmethod
    def render_chart(
        slide: PptxSlide,
        chart_block: ChartBlock,
        layout_result: ChartLayoutResult,
        x: float,
        y: float,
        width: float,
        height: float
    ):
        """
        Renders a native PowerPoint chart using CategoryChartData.
        """
        chart_data = CategoryChartData()
        chart_data.categories = chart_block.categories

        for series in chart_block.series:
            chart_data.add_series(series.name, series.values)

        c_type = chart_block.chart_type.lower()
        if c_type == "pie":
            xl_type = XL_CHART_TYPE.PIE
        elif c_type == "line":
            xl_type = XL_CHART_TYPE.LINE
        else:
            xl_type = XL_CHART_TYPE.COLUMN_CLUSTERED

        left, top, w, h = bbox_to_pptx_emu(x, y, width, height)
        chart_shape = slide.shapes.add_chart(xl_type, left, top, w, h, chart_data)
        chart = chart_shape.chart

        # Title
        if chart_block.title:
            chart.has_title = True
            chart.chart_title.text_frame.text = chart_block.title
            for p in chart.chart_title.text_frame.paragraphs:
                p.font.name = Theme.FONT_PRIMARY
                p.font.size = Pt(12)
                p.font.bold = True
                p.font.color.rgb = hex_to_rgb_color(Theme.PRIMARY_NAVY)
        else:
            chart.has_title = False

        # Legend
        if layout_result.legend_position != "none":
            chart.has_legend = True
            legend_map = {
                "top": XL_LEGEND_POSITION.TOP,
                "bottom": XL_LEGEND_POSITION.BOTTOM,
                "right": XL_LEGEND_POSITION.RIGHT,
                "left": XL_LEGEND_POSITION.LEFT
            }
            chart.legend.position = legend_map.get(layout_result.legend_position, XL_LEGEND_POSITION.RIGHT)
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(9)
        else:
            chart.has_legend = False

        # Style series with theme colors
        palette = Theme.CHART_PALETTE
        try:
            if c_type == "pie" and len(chart.plots) > 0 and len(chart.plots[0].series) > 0:
                # Color pie slices
                series = chart.plots[0].series[0]
                for idx, point in enumerate(series.points):
                    color_hex = palette[idx % len(palette)]
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = hex_to_rgb_color(color_hex)
            else:
                for idx, series in enumerate(chart.series):
                    color_hex = palette[idx % len(palette)]
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = hex_to_rgb_color(color_hex)
        except Exception:
            # Fallback gracefully if specific color application isn't supported on plot type
            pass

        return chart_shape
