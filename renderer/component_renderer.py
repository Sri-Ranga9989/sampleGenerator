"""
Shared Component Renderer
Renders reusable visual components (insight cards, callout badges, container borders).
"""

from pptx.slide import Slide as PptxSlide
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from layout.coordinate_system import bbox_to_pptx_emu
from config.theme import Theme, hex_to_rgb_color
from models.report_model import InsightBlock


class ComponentRenderer:
    @staticmethod
    def render_insight_card(
        slide: PptxSlide,
        insight: InsightBlock,
        x: float,
        y: float,
        width: float,
        height: float,
        accent_color: str = Theme.ACCENT_CYAN
    ):
        """
        Renders an analytical insight callout card:
        - Card fill: #F8FAFC
        - Left accent bar: accent_color (e.g. #3894D8)
        - Title: bold #0D3166
        - Body: #1F2937
        """
        # 1. Card background shape
        left, top, w, h = bbox_to_pptx_emu(x, y, width, height)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = hex_to_rgb_color(Theme.BG_CARD)
        card.line.color.rgb = hex_to_rgb_color(Theme.BORDER_LIGHT)
        card.line.width = Pt(1)

        # 2. Left accent bar
        bar_w = 6.0
        bar_left, bar_top, b_w, b_h = bbox_to_pptx_emu(x + 2, y + 2, bar_w, height - 4)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, bar_left, bar_top, b_w, b_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = hex_to_rgb_color(accent_color)
        bar.line.fill.background()

        # 3. Card text
        text_x = x + bar_w + 14.0
        text_w = max(20.0, width - bar_w - 24.0)
        t_left, t_top, t_w, t_h = bbox_to_pptx_emu(text_x, y + 6.0, text_w, height - 12.0)
        tx_box = slide.shapes.add_textbox(t_left, t_top, t_w, t_h)
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_top = 0

        # Title
        p0 = tf.paragraphs[0]
        p0.text = f"KEY INSIGHT: {insight.title}" if not insight.title.upper().startswith("KEY") else insight.title
        p0.font.name = Theme.FONT_PRIMARY
        p0.font.size = Pt(10.5)
        p0.font.bold = True
        p0.font.color.rgb = hex_to_rgb_color(Theme.PRIMARY_NAVY)
        p0.space_after = Pt(4)

        # Body
        p1 = tf.add_paragraph()
        p1.text = insight.body
        p1.font.name = Theme.FONT_PRIMARY
        p1.font.size = Pt(9.5)
        p1.font.bold = False
        p1.font.color.rgb = hex_to_rgb_color(Theme.TEXT_DARK)

        return card, bar, tx_box
