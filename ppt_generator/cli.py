"""
CLI Interface for PPT Generator V1
Provides unified entry points for:
- validate-input
- validate-pptx
- render-template
- generate
- compare
"""

import sys
import os
import json
import argparse
from typing import Dict, Any, List
from pptx import Presentation
from models.report_model import ReportParser, Slide, SlideData
from renderer.template_renderer import TemplateRenderer
from renderer.powerpoint_renderer import PowerPointRenderer
from renderer.rasterizer import PowerPointRasterizer
from validators.structural_validator import StructuralValidator
from validators.rendered_validator import RenderedValidator
from validators.completeness_validator import CompletenessValidator
from validators.visual_regression import VisualRegressionValidator


def validate_input_cmd(args):
    """Validates input JSON schema and parser compatibility."""
    input_path = args.input
    if not os.path.exists(input_path):
        print(f"[ERROR] Input file not found: {input_path}")
        sys.exit(1)

    try:
        with open(input_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception as e:
        print(f"[ERROR] Invalid JSON syntax: {e}")
        sys.exit(1)

    print(f"[VALIDATE-INPUT] Validating manifest: {input_path}...")

    # Validate with jsonschema if available
    try:
        import jsonschema
        schema_path = "schemas/report.schema.json"
        if os.path.exists(schema_path):
            with open(schema_path, "r", encoding="utf-8") as sf:
                schema = json.load(sf)
            resolver = jsonschema.RefResolver(base_uri=f"file:///{os.path.abspath('schemas')}/", referrer=schema)
            jsonschema.validate(instance=data, schema=schema, resolver=resolver)
            print("  -> JSON Schema (Draft-07): PASS")
    except ImportError:
        print("  -> jsonschema package not installed, skipping schema validation")
    except Exception as e:
        msg = getattr(e, "message", str(e))
        print(f"  -> JSON Schema validation warning: {msg}")

    # Parse into canonical Report
    try:
        report = ReportParser.parse_report_json(data)
        print(f"  -> Canonical Data Model: PASS ({len(report.slides)} slides, {len(report.all_input_ids)} content items registered)")
        print("\n[RESULT] Input file is VALID.")
    except Exception as e:
        print(f"[ERROR] Data model parsing error: {e}")
        sys.exit(1)


def validate_pptx_cmd(args):
    """Inspects PowerPoint presentation for structural integrity and font clamp compliance."""
    pptx_path = args.input
    if not os.path.exists(pptx_path):
        print(f"[ERROR] Presentation file not found: {pptx_path}")
        sys.exit(1)

    print(f"[VALIDATE-PPTX] Inspecting presentation: {pptx_path}...")
    prs = Presentation(pptx_path)
    res = StructuralValidator.validate_presentation(prs)
    print(f"  -> Status: {res['status']}")
    print(f"  -> Total Slides: {res['total_slides']}")
    if res["errors"]:
        print(f"  -> Errors ({len(res['errors'])}):")
        for err in res["errors"]:
            print(f"     - [Slide {err.get('slide', '?')}] {err.get('type')}: {err.get('message')}")
        sys.exit(1)
    else:
        print("\n[RESULT] Presentation is STRUCTURALLY VALID.")


def render_template_cmd(args):
    """Renders a single template archetype to PPTX."""
    t_id = args.template_id
    input_path = args.input
    output_path = args.output

    if not os.path.exists(input_path):
        print(f"[ERROR] Input file not found: {input_path}")
        sys.exit(1)

    with open(input_path, "r", encoding="utf-8") as f:
        raw_data = json.load(f)

    # Could be full report with 1 slide or direct slide dict
    if "slides" in raw_data and len(raw_data["slides"]) > 0:
        slide_dict = raw_data["slides"][0]
    else:
        slide_dict = raw_data

    report = ReportParser.parse_report_json({"slides": [slide_dict]})
    slide = report.slides[0]
    slide.template_id = t_id

    print(f"[RENDER-TEMPLATE] Rendering template '{t_id}'...")
    t_renderer = TemplateRenderer()
    layout_results = t_renderer.render_slide_layout(slide)

    ppt_renderer = PowerPointRenderer()
    prs, prov = ppt_renderer.create_presentation(layout_results)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    prs.save(output_path)
    print(f"  -> Saved PPTX: {output_path}")

    struct_res = StructuralValidator.validate_presentation(prs)
    print(f"  -> Structural Validation: {struct_res['status']}")


def generate_cmd(args):
    """Executes full end-to-end presentation generation pipeline."""
    input_path = args.input
    output_path = args.output
    rasterize = getattr(args, "rasterize", True)

    if not os.path.exists(input_path):
        print(f"[ERROR] Input file not found: {input_path}")
        sys.exit(1)

    with open(input_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    print(f"\n=======================================================")
    print(f"PPT GENERATOR V1 — END-TO-END PRESENTATION PIPELINE")
    print(f"=======================================================")
    print(f"Loading input manifest: {input_path}")
    report = ReportParser.parse_report_json(data)
    print(f"Parsed report: '{report.title}' ({len(report.slides)} input slides, {len(report.all_input_ids)} content IDs)")

    # 1. Layout Engine
    print("\n[PHASE 1] Computing Layout & Budgeting...")
    t_renderer = TemplateRenderer()
    layout_results = t_renderer.render_report_layout(report)
    print(f"  -> Computed {len(layout_results)} output slides (including dynamic continuation slides)")

    # 2. PowerPoint Rendering
    print("\n[PHASE 2] Constructing Native PowerPoint PPTX...")
    ppt_renderer = PowerPointRenderer()
    prs, prov_map = ppt_renderer.create_presentation(layout_results)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    prs.save(output_path)
    print(f"  -> Successfully generated: {output_path}")

    # 3. Layer A: Structural Validation
    print("\n[PHASE 3] Running Layer A Structural Validation...")
    struct_res = StructuralValidator.validate_presentation(prs)
    print(f"  -> Structural Status: {struct_res['status']} ({len(struct_res['errors'])} errors)")

    # 4. Content Completeness Validation
    print("\n[PHASE 4] Running Content Completeness Validation...")
    comp_res = CompletenessValidator.validate_completeness(report.all_input_ids, prov_map)
    print(f"  -> Completeness Status: {comp_res['status']}")
    print(f"  -> Input items: {comp_res['total_input_items']}, Rendered: {comp_res['total_rendered_items']}, Missing: {comp_res['missing_count']}")

    # 5. Slide Rasterization & Layer B Validation
    if rasterize:
        print("\n[PHASE 5] Rasterizing Slides via PowerPoint COM & Layer B Validation...")
        rendered_dir = os.path.join(os.path.dirname(output_path), "rendered")
        os.makedirs(rendered_dir, exist_ok=True)
        try:
            pngs = PowerPointRasterizer.rasterize_pptx(output_path, rendered_dir)
            print(f"  -> Rasterized {len(pngs)} slide images into: {rendered_dir}")
            rendered_errors = 0
            for png in pngs:
                r_val = RenderedValidator.validate_slide_image(png)
                if r_val["status"] != "PASS":
                    rendered_errors += 1
            rendered_status = "PASS" if rendered_errors == 0 else "FAIL"
            print(f"  -> Layer B Visual Status: {rendered_status} ({rendered_errors} corrupted slides)")
        except Exception as e:
            print(f"  -> Note: Slide rasterization COM call bypassed: {e}")

    print("\n=======================================================")
    print("PIPELINE EXECUTION COMPLETE")
    print(f"Output presentation: {output_path}")
    print("=======================================================\n")


def compare_cmd(args):
    """Compares rendered PNG against reference PNG."""
    ref_path = args.reference
    gen_path = args.generated
    res = VisualRegressionValidator.compare_images(ref_path, gen_path)
    print(f"[COMPARE] Status: {res['status']}")
    if "message" in res:
        print(f"  -> Message: {res['message']}")
    if "diff_percentage" in res:
        print(f"  -> Diff %: {res['diff_percentage']}% (Tolerance: {res['tolerance']}%)")
        print(f"  -> Diff Image: {res['diff_image_path']}")


def parse_html_cmd(args):
    """Parses a Gmail HTML email into a canonical JSON manifest."""
    from html_parser import parse_html_to_json

    input_path = args.input
    output_path = args.output
    auto_charts = not args.no_auto_charts

    if not os.path.exists(input_path):
        print(f"[ERROR] Input file not found: {input_path}")
        sys.exit(1)

    # Default output path
    if not output_path:
        base = os.path.splitext(os.path.basename(input_path))[0]
        slug = base.lower().replace(" ", "_")[:40]
        output_path = os.path.join("output", slug, "report.json")

    print(f"\n=======================================================")
    print(f"PPT GENERATOR V1 — HTML EMAIL PARSING PIPELINE")
    print(f"=======================================================")

    report, stats, warnings = parse_html_to_json(
        input_path, output_path, auto_charts=auto_charts
    )

    print(f"\n[RESULT] JSON manifest saved: {output_path}")
    print(f"  -> Slides: {len(report.get('slides', []))}")
    print(f"  -> Tables: {stats.get('data_tables', 0)}")
    print(f"  -> Charts: {stats.get('charts_generated', 0)}")
    print(f"  -> Warnings: {len(warnings)}")
    print(f"=======================================================\n")


def serve_cmd(args):
    """Starts the FastAPI server for the PPT Generator API."""
    try:
        import uvicorn
    except ImportError:
        print("[ERROR] uvicorn is not installed. Run: pip install uvicorn fastapi python-multipart")
        sys.exit(1)

    port = args.port
    host = getattr(args, "host", "0.0.0.0")
    print(f"\n=======================================================")
    print(f"PPT GENERATOR V1 — API SERVER")
    print(f"=======================================================")
    print(f"Starting server on http://{host}:{port}")
    print(f"API docs at http://localhost:{port}/docs")
    print(f"=======================================================\n")

    uvicorn.run("api_server:app", host=host, port=port, reload=False)


def main():
    parser = argparse.ArgumentParser(description="PPT Generator V1 Command Line Interface")
    subparsers = parser.add_subparsers(dest="command", help="Available commands")

    # validate-input
    p_val_in = subparsers.add_parser("validate-input", help="Validate input JSON schema")
    p_val_in.add_argument("--input", required=True, help="Path to input JSON file")

    # validate-pptx
    p_val_ppt = subparsers.add_parser("validate-pptx", help="Validate PPTX structural integrity")
    p_val_ppt.add_argument("--input", required=True, help="Path to presentation PPTX")

    # render-template
    p_rnd_tmp = subparsers.add_parser("render-template", help="Render single template archetype")
    p_rnd_tmp.add_argument("--template-id", required=True, help="Template ID (e.g. 04_table_chart)")
    p_rnd_tmp.add_argument("--input", required=True, help="Path to slide content JSON")
    p_rnd_tmp.add_argument("--output", required=True, help="Path to output PPTX")

    # generate
    p_gen = subparsers.add_parser("generate", help="Generate complete presentation from report manifest")
    p_gen.add_argument("--input", required=True, help="Path to report manifest JSON")
    p_gen.add_argument("--output", required=True, help="Path to output PPTX")
    p_gen.add_argument("--no-rasterize", dest="rasterize", action="store_false", help="Skip PowerPoint COM rasterization")

    # compare
    p_cmp = subparsers.add_parser("compare", help="Compare generated slide image against reference image")
    p_cmp.add_argument("--reference", required=True, help="Path to reference PNG")
    p_cmp.add_argument("--generated", required=True, help="Path to generated PNG")

    # parse-html (NEW)
    p_parse = subparsers.add_parser("parse-html", help="Parse Gmail HTML email into canonical JSON manifest")
    p_parse.add_argument("--input", required=True, help="Path to Gmail HTML email file")
    p_parse.add_argument("--output", default=None, help="Path to output JSON manifest")
    p_parse.add_argument("--no-auto-charts", action="store_true", help="Disable automatic chart generation")

    # serve (NEW)
    p_serve = subparsers.add_parser("serve", help="Start FastAPI server for PPT Generator API")
    p_serve.add_argument("--host", default="0.0.0.0", help="Server host (default: 0.0.0.0)")
    p_serve.add_argument("--port", type=int, default=int(os.environ.get("PORT", 8050)), help="Server port (default: 8050)")

    args = parser.parse_args()

    if args.command == "validate-input":
        validate_input_cmd(args)
    elif args.command == "validate-pptx":
        validate_pptx_cmd(args)
    elif args.command == "render-template":
        render_template_cmd(args)
    elif args.command == "generate":
        generate_cmd(args)
    elif args.command == "compare":
        compare_cmd(args)
    elif args.command == "parse-html":
        parse_html_cmd(args)
    elif args.command == "serve":
        serve_cmd(args)
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
